from pathlib import Path
import json, copy
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[2]; out=Path(__file__).resolve().parent.parent/'实验结果'; report=Path(__file__).resolve().parent.parent/'报告'
info=json.loads((out/'实验元数据.json').read_text());ss=info['stats'];s=ss[0]
by={(v['mode'],v['threads']):v for v in ss}
a4=by['static',4];l4=by['lock',4];o4=by['atomic',4]
labels={'serial':'串行','static':'静态分块','lock':'锁任务池','atomic':'原子任务池'}
fmt=lambda x:f'{x:.3f}'
template=root/'temp/报告原始模版.docx'
d=Document(template) if template.exists() else Document()
# Reuse the supplied document's page geometry and styles, replace all example content.
# Clear legacy compatibility/grid settings that cause clipped CJK text in conversion.
fresh=Document()
settings=d.settings._element
for child in list(settings): settings.remove(child)
for child in fresh.settings._element: settings.append(copy.deepcopy(child))
body=d._element.body
for child in list(body):
 if child.tag!=qn('w:sectPr'):body.remove(child)
sect=d.sections[0]
cols=sect._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1')
grid=sect._sectPr.find(qn('w:docGrid'))
if grid is not None: sect._sectPr.remove(grid)
sect.top_margin=Cm(1.8);sect.bottom_margin=Cm(1.8)
sect.left_margin=sect.right_margin=Cm(1.7)
for header in [sect.header,sect.footer]:
 for el in list(header._element):header._element.remove(el)
p=sect.footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

def font(run,name='宋体',size=10.5,bold=False):
 run.font.name='Noto Serif CJK SC' if name=='宋体' else 'Noto Sans CJK SC' if name=='黑体' else 'Noto Serif CJK SC'
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),run.font.name)
 run.font.size=Pt(size);run.bold=bold
for name in ['Normal','Body Text']:
 st=d.styles[name];st.font.size=Pt(10.5);st.font.name='Noto Serif CJK SC'
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 st.paragraph_format.line_spacing=Pt(14);st.paragraph_format.space_after=Pt(3)

def para(text='',indent=True,align=None,size=10.5,name='宋体'):
 p=d.add_paragraph();p.paragraph_format.first_line_indent=Pt(21) if indent else Pt(0)
 p.paragraph_format.widow_control=True
 p.paragraph_format.line_spacing=Pt(max(14,size*1.35))
 if align is not None:p.alignment=align
 font(p.add_run(text),name,size)
 return p

def heading(text,level=1):
 p=d.add_paragraph();p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
 font(p.add_run(text),'黑体',12 if level==1 else 10.5,True)
 return p

def section(columns=1,newpage=True):
 sec=d.add_section(WD_SECTION_START.NEW_PAGE if newpage else WD_SECTION_START.CONTINUOUS)
 sec._sectPr.find(qn('w:cols')).set(qn('w:num'),str(columns))
 sec._sectPr.find(qn('w:cols')).set(qn('w:space'),'425')
 return sec

def caption(zh,en):
 p=para(zh,False,WD_ALIGN_PARAGRAPH.CENTER,9);p.paragraph_format.keep_with_next=True
 p=para(en,False,WD_ALIGN_PARAGRAPH.CENTER,9)

def table(headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=True
 for cells,vals in [(t.rows[0].cells,headers)]+[(t.add_row().cells,r) for r in rows]:
  for c,v in zip(cells,vals):
   c.text='';p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=Pt(12)
   font(p.add_run(str(v)),'宋体',9)
 # Three-line table, no vertical rules.
 for ri,row in enumerate(t.rows):
  no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
  for c in row.cells:
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);on=(edge=='top' and ri==0) or(edge=='bottom' and (ri==0 or ri==len(t.rows)-1))
    e.set(qn('w:val'),'single' if on else 'nil');e.set(qn('w:sz'),'8' if ri in [0,len(t.rows)-1] else '4');borders.append(e)
   c._tc.get_or_add_tcPr().append(borders)
 return t

def pic(filename,width=16):
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.line_spacing=1.0
 p.add_run().add_picture(str(out/filename),width=Cm(width))
 p.paragraph_format.space_after=Pt(2)
 p.paragraph_format.keep_with_next=True

para('基于 OpenMP 的 Sobel 图像边缘检测\n任务划分、同步机制与并行性能分析',False,WD_ALIGN_PARAGRAPH.CENTER,16)
para('学号：________    姓名：________',False,WD_ALIGN_PARAGRAPH.CENTER,14)
para('（学院、专业及班级：________________）',False,WD_ALIGN_PARAGRAPH.CENTER,9)
para('实验日期：2026年10月25日',False,WD_ALIGN_PARAGRAPH.CENTER,9)
para(f'摘要  针对Sobel边缘检测中大量独立的邻域计算，本文从任务分解与共享数据依赖出发，设计串行、静态分块、锁任务池和原子任务池四种实现。各版本采用相同卷积公式及输出规则，将固定的4096×4096灰度图划分为64个行块，保持每块最多64行。静态方案预先分配任务，动态方案分别用锁和原子捕获保护任务编号，像素计算均在临界区外进行。实验固定输入与编译参数，比较4线程下的不同设计，以及1、2、4线程下的扩展性。每组预热3次、正式测量30次，并在独立诊断运行中记录任务分配与线程工作时间。串行平均耗时为{s["mean_ms"]:.3f} ms，三种4线程方案分别为{a4["mean_ms"]:.3f}、{l4["mean_ms"]:.3f}和{o4["mean_ms"]:.3f} ms。当前云环境下，动态任务分配优于静态分块，但结果不能直接解释为锁操作优于原子操作。本文结合依赖关系、任务所有权与实测诊断讨论并行收益及其适用范围。',False,name='楷体')
para('关键词  Sobel；OpenMP；任务分解；共享任务池；互斥锁；原子操作；并行效率',False,name='楷体')
section(2,False)
heading('1 引言')
para('边缘检测根据局部灰度变化提取轮廓，是图像处理的基础操作。Sobel算子以两个3×3卷积核估计水平和垂直梯度。固定图像中，各输出像素只依赖原图邻域，不依赖其他输出值，因而具有明确的数据并行性。')
para('本作业围绕“针对性地用并行程序解决问题”展开[1]。研究重点不是启动多少线程，而是如何划分工作、安排数据访问、保护确实存在的共享状态，以及如何通过实验判断这些选择是否有效。本文只实现Sobel，不增加其他图像处理算法；性能实验限定为设计方案对比与线程扩展两项。')
heading('2 问题分析与总体设计')
heading('2.1 从数据依赖确定并行边界',2)
para('设原图为I、输出为O。O(x,y)只读取I中以(x,y)为中心的九个像素。输入从读取完成后保持不变，不同任务负责互不重叠的输出行。因此，相邻任务即使读取同一输入行，也没有读写冲突；无需给像素读取或计算加锁。')
para('必须区分“无冲突的输出计算”和“有冲突的任务分配”。静态方案没有共享领取状态；任务池方案需要多个线程共同更新下一个任务编号next，这才是互斥或原子操作应该保护的对象。所有方案都在并行区域结束时汇合，避免主线程提前读取未完成结果。')
heading('2.2 用PCAM组织设计',2)
para('课程中的PCAM将设计分为划分、通信、组合和映射[2]。本文据此作出以下选择。')
para('划分：以输出像素为独立工作，沿行方向建立任务。通信：输入只读共享，任务读取真实邻行，不复制或交换整幅图像。组合：将连续像素合并成行块，避免逐像素领取任务。映射：比较预先均分任务与运行时按需领取。')
para('选择连续行块是因为图像按行连续存储，便于顺序访问；每块最多64行，在4094个内部有效行上形成64个任务，4线程时平均每线程16块，既保留重新分配的余地，又避免每个像素都访问同步变量。64行是预先固定的实验参数，不声称它是最优粒度。')
heading('2.3 程序模块与数据流',2)
para('I/O模块读取8位P5 PGM；串行与并行模块执行共同的单行Sobel公式；验证模块使用独立参考实现；基准模块负责计时与结果校验；Python脚本汇总CSV和绘图。图像分配、文件读写、结果比较和绘图均与核心计时分离。')
section(1)
heading('3 任务划分与并行实现')
pic('设计_处理流程.png',14)
caption('图1 程序流程与控制变量','Fig. 1 Program workflow and controlled variables')
heading('3.1 一致的Sobel计算规则',2)
para('水平与垂直卷积核分别为Kx=[−1,0,1;−2,0,2;−1,0,1]、Ky=[−1,−2,−1;0,0,0;1,2,1][5]。内部像素输出为floor(min(255, √(Gx²+Gy²)))，不做额外归一化或二值化。梯度以int累加，平方和转double后开方；最外圈输出直接为0。')
para('每次调用覆盖全部输出，包括边界；输入输出使用不同缓冲区。四种实现复用相同的单行函数，串行版也顺序遍历同一组64行任务块。这样，方案间的工作量和输出定义保持一致，差异集中于任务映射和同步。')
heading('3.2 行块定义与接缝处理',2)
para('令H为图像高度，B=64。内部行数N=max(0,H−2)，任务数K=ceil(N/B)。任务k的起始行为1+kB，结束行为min(H−1,1+(k+1)B)，采用左闭右开区间。4096行输入对应N=4094、K=64：前63块各64行，最后一块62行。')
para('任务k读取自身输出区域上下各一行真实输入。相邻块的输入读取可以重叠，输出写入不能重叠。任务接缝不是图像边界，不能在接缝处置零或截断邻域，否则即使各线程互不冲突，也会生成错误的边缘。')
pic('设计_任务边界.png',12)
caption('图2 相邻任务的输出所有权与输入邻行','Fig. 2 Output ownership and neighboring input rows')
heading('3.3 四个版本的设计职责',2)
table(['版本','任务映射','共享状态与同步','设计目的'],[
['S 串行','按编号依次处理全部块','无工作线程','提供真正的串行基准'],
['P1 静态','启动时计算连续任务区间','结束时隐式屏障','低分配成本，显式任务所有权'],
['P2 锁任务池','完成一块后继续领取','锁保护next；结束时屏障','应对实际执行速率差异'],
['P3 原子任务池','与P2相同的按需领取','atomic capture更新next','比较单整数共享状态的实现选择']])
section(2)
heading('3.4 静态映射的设计与正确性',2)
para('在OpenMP parallel区域中，线程t依据实际团队大小p取得任务区间[floor(Kt/p),floor(K(t+1)/p))。相邻区间首尾相接且互不重叠，所有区间的并集恰好是[0,K)，所以不存在遗漏或重复。这也能处理K不能被p整除、甚至K小于p的情况。')
para('静态方案只在启动时计算范围，不必每块访问共享调度变量。对于计算量均匀、线程执行速率近似一致的环境，这是合理的低开销设计。但均分任务数并不保证同时完成：相同工作量的线程仍可能受到不同的实际调度、频率或资源竞争影响。')
heading('3.5 锁任务池：只保护共享领取状态',2)
para('任务池不存放大量对象，只用共享整数next表示尚未领取的首个任务。每次核心调用将next置0并初始化一把OpenMP简单锁。线程循环执行以下过程：')
para('加锁 → task=next并递增next → 解锁\n若task≥K则退出\n否则在锁外计算该任务，再次领取',False,size=9)
para('如果不保护“读取并递增”，多个线程可能同时读取同一编号，造成重复任务和遗漏任务。互斥将领取操作串行化，保证成功领取的编号唯一。线程拿到编号后立即释放锁，整个行块的Sobel计算不会占有该锁。')
para('线程仅在持锁时更新next，不在锁内等待其他线程、不做I/O，也不嵌套申请其他锁。并行区域汇合后销毁锁，确保没有线程仍使用它。这种结构避免了循环持锁等待；它不意味着任意修改后的程序都不可能死锁。')
para('动态方案的机会在于较早完成的线程可以继续领取任务，减少固定分配造成的尾部等待。代价是每块领取与锁竞争。全图有K次成功领取，每个线程另有一次越界领取用于退出，因此总领取次数为K+p；4线程下为68次，而不是每像素一次。')
heading('3.6 原子任务池：适用条件与限制',2)
para('P3保留P2的任务池逻辑，只把“加锁—取值递增—解锁”替换成OpenMP atomic capture，即以一个原子操作取得旧编号并递增共享整数[3]。像素计算仍在操作之外，所有任务完成后仍需并行区域末尾的同步。')
para('这种替换成立的前提是共享领取状态仅为一个整数。若真实队列还需同时维护链表、队列头尾或任务状态，单个原子递增不能自动保证整个复合操作正确。原子操作也会竞争同一个共享位置，不能先验断言其整体程序必然比锁版本更快。')
heading('3.7 共享变量、局部变量与屏障',2)
para('输入和输出对象共享，next及锁在任务池中共享；线程编号、块编号、循环索引和梯度变量均为线程局部。使用default(none)显式约束共享关系。各任务无需等待其他任务的中间结果，不设置块间屏障；并行区域末尾的隐式屏障则不可省略。')
para('这体现了互斥与汇合的不同职责：锁或原子操作解决“谁领取哪个任务”，屏障解决“所有结果何时可被统一使用”。不在像素计算中加锁，是依据数据依赖作出的设计决定，而不是遗漏同步。')
heading('3.8 分析模型与适用范围',2)
para('各版本计算复杂度均为O(WH)，主要图像缓冲区空间为O(WH)。动态分配额外执行O(K+p)次领取。可用Tp≈T边界+T最慢线程工作+T分配与汇合描述实际时间，并结合含并行开销的Amdahl模型理解加速上限[4]。')
para('增大并行度可能缩短像素计算，却不会消除分配、同步与资源争用。本实验不拟合串行比例，也不把所有非理想加速都归因于锁；这些因素需要依靠实测现象和适用范围分别讨论。')
section(1)
heading('4 实验设计：测什么，为什么这样测')
heading('4.1 实验环境与固定输入',2)
caption('表1 实验环境与不变条件','Table 1 Environment and fixed conditions')
table(['项目','配置'],[
['处理器与运行环境','Intel Xeon Platinum 8573C；KVM云端虚拟环境'],
['可见CPU与配额','5个逻辑CPU；cpu.max=400000 100000，等效4核'],
['系统与内存限制','Debian GNU/Linux 13，x86_64；32 GiB'],
['编译器与运行时','G++ 14.2.0；libgomp；_OPENMP=201511'],
['统一编译选项','-O3 -std=c++17 -fopenmp -Wall -Wextra -Wpedantic'],
['输入与任务粒度','4096×4096，8位灰度；64行/块，64个任务'],
['OpenMP配置','DYNAMIC=FALSE；PROC_BIND=close；PLACES=cores；WAIT_POLICY=PASSIVE'],
['测量重复与顺序','每配置预热3次；正式30轮；每轮打乱10种配置'],
['线程数与实验范围','1、2、4线程；只比较方案和线程数，不改变图像尺寸']])
para('固定输入由脚本生成，包含线性灰度背景、矩形、圆、三角形和条纹，不是自然图像数据集。它提供可视轮廓与可复现输入；Sobel对每个内部像素执行相同公式，不依据边缘内容跳过计算。所有组共用一份PGM，文件随提交保存。')
para('输入PGM的SHA-256：'+info['input_sha256'],False,size=9)
heading('4.2 两项实验的自变量与控制变量',2)
caption('表2 实验矩阵与设计理由','Table 2 Experimental matrix and rationale')
table(['实验','变化因素','比较对象','需要回答的问题'],[
['一：设计方案','分配与同步方案','S，以及P1/P2/P3各4线程','哪种设计在本环境下更有效？'],
['二：线程扩展','线程数；分别观察各方案','P1/P2/P3各1、2、4线程','扩展趋势是否取决于分配方式？']])
para('固定图像尺寸可把性能变化与工作规模变化分开，属于固定规模的强扩展考察。固定64行任务块可避免把任务粒度与同步方式同时改变；不比较动态调度的不同chunk参数，不增加其他图像算法。4线程由有效CPU配额预先确定，不把超出配额的8线程作为主要测试点。')
para('实验一比较三种并行设计，而不只比较“有没有线程”；实验二在每个设计内部观察线程增长。两项实验共享同一次测量中的S、P1/P2/P3数据，没有为了某个方案更换输入、编译参数或基准。')
para('预期并不作为结论：均匀工作可能使静态分配占优，运行速率差异也可能使任务池占优。原子领取的临界操作更专门化，但整体耗时是否更短仍需测量。允许结果否定先验判断，才有比较设计方案的意义。')
section(2)
heading('4.3 正式计时与统计设计',2)
para('共10种配置：1组串行与3方案×3线程数。程序先验证实际线程团队可达到1、2、4，再分别预热3次。正式阶段进行30轮，每轮用固定种子20261025打乱10种配置的顺序，共得到300条样本。打乱顺序用于减轻时间漂移对固定先后顺序的偏置，不保证彻底消除云端干扰。')
para('所有配置使用同一个主线程上的omp_get_wtime计时[6]，计时包围完整核心调用及共同的配置分派。包括边界写入、实际像素运算、线程进入与退出、任务领取、锁初始化与销毁及汇合同步；排除输入输出、缓冲区分配、结果比较、统计和绘图。')
para('预热可以减轻首次页访问和首次线程启动的影响，因此本报告描述预热后的重复运行表现，不把结果解释为首次运行或端到端图像处理性能。输出缓冲区每次均被完整覆盖，不能用上次结果代替本次计算。')
para('每组取30次算术均值，标准差使用n−1分母；保留全部样本并附中位数与极值。图中误差线是标准差，不是置信区间，不以误差线重叠与否代替显著性检验。')
para('加速比Sp=Tserial/Tparallel(p)，效率Ep=Sp/p。统一使用独立串行实现的均值作分母，而非将OpenMP的1线程结果当作串行基准，符合课程的性能评价定义[4]。')
heading('4.4 诊断运行为什么独立进行',2)
para('完成正式计时后，另对9个并行配置各诊断3次，共27次。记录每线程完成的块数、内部行数、领取次数以及进入工作阶段到完成领取循环的耗时。后者包含线程被调度出去的时间，不包含末尾屏障等待，也不是纯CPU计算时间。')
para('任务访问计数使用原子更新，每个任务应恰好访问一次。诊断逻辑以编译期模板参数区分，正式版本不执行计数器更新和每线程时钟调用。这样避免诊断额外引入的同步成本污染主结果。诊断只用于解释现象，不能替代正式耗时。')
para('对4线程，静态版本预期每线程16块，而任务池允许分配不均。这里的“不均”可能正是对线程实际执行速率差异的适应，不能仅凭任务数不相等就判定调度差。线程耗时差也不直接证明某个具体硬件瓶颈。')
heading('4.5 正确性验收与边界用例',2)
para('验证在性能结论之前进行。独立参考实现显式遍历两个3×3系数矩阵，不调用待测单行函数。覆盖11种尺寸、5类输入：全零、常数、水平斜坡、阶跃和固定种子随机图。高度66、67、131、257覆盖64行整块、尾块及任务分界。')
para('每幅图先验证串行，再组合3方案×4种线程数（1、2、3、4）×4种块高（1、3、16、64），共2695次参考比较。这里改变块高只用于边界正确性验收，不加入性能实验。另执行2640次诊断所有权检查，验证任务唯一、行数正确及领取次数符合K+p。')
para('手算5×5斜坡图的中心梯度为Gx=24、Gy=0，输出24；另检查非法线程数、非法块高及输入输出别名被拒绝。普通构建和ASan/UBSan构建均通过。此类内存与未定义行为检查不能替代数据竞争检测；并发正确性仍依赖所有权分析、领取原子性及验收结果。')
para('固定全图的300次正式输出与串行结果逐像素一致，27次诊断输出也一致，且每个任务恰好访问一次。所有输出均为4096×4096，最外圈为0。图像展示只是直观补充，不用肉眼相似代替数值一致。')
section(1)
heading('5 实验结果与机制分析')
heading('5.1 实验一：固定4线程的设计方案比较',2)
caption('表3 串行与三种4线程实现的正式测量','Table 3 Formal measurements of serial and four-thread implementations')
table(['配置','均值/ms','标准差/ms','加速比','效率/%'],[
[labels[v['mode']],fmt(v['mean_ms']),fmt(v['std_ms']),fmt(v['speedup']),'—' if v['mode']=='serial' else f'{v["efficiency"]*100:.2f}'] for v in [s,a4,l4,o4]])
pic('实验一_方案对比.png',15)
caption('图3 固定4线程时的方案比较（均值±样本标准差）','Fig. 3 Design comparison at four threads (mean ± sample SD)')
para(f'三种方案相对串行均有平均加速。静态、锁和原子任务池的加速比分别为{a4["speedup"]:.3f}、{l4["speedup"]:.3f}和{o4["speedup"]:.3f}。锁任务池在本次测量中均值最小，相对静态分块缩短{(1-l4["mean_ms"]/a4["mean_ms"])*100:.2f}%，原子任务池相对静态缩短{(1-o4["mean_ms"]/a4["mean_ms"])*100:.2f}%。')
para('这一结果说明，在当前环境和固定64行粒度下，按需领取任务所获得的收益超过了额外领取成本。但不能据此得出“锁比无锁计算快”：静态方案和任务池还改变了任务映射方式。更不能由两个程序的总时间差直接计算单次加锁或单次原子操作成本。')
para(f'锁任务池比原子任务池的平均时间低{o4["mean_ms"]-l4["mean_ms"]:.3f} ms，而两组标准差分别为{l4["std_ms"]:.3f}和{o4["std_ms"]:.3f} ms。这里只报告本次整体实现的测量排序，不将其提升为锁原语性能的普遍结论。原子操作仍有共享位置争用，领取节奏、任务落点和云端波动都可能影响整体结果。')
section(1)
heading('5.2 实验二：线程扩展与效率',2)
caption('表4 各设计在1、2、4线程下的性能','Table 4 Performance of each design at one, two and four threads')
table(['方案','线程','均值/ms','标准差/ms','加速比','效率/%'],[
[labels[v['mode']],v['threads'],fmt(v['mean_ms']),fmt(v['std_ms']),fmt(v['speedup']),f'{v["efficiency"]*100:.2f}'] for v in ss[1:]])
pic('实验二_线程扩展.png',17)
caption('图4 线程数量对核心时间、加速比与效率的影响','Fig. 4 Thread scaling of runtime, speedup and efficiency')
para('三种设计从1线程增加到4线程，平均时间均下降，但扩展趋势不同。静态方案从2线程到4线程只进一步减少'+f'{(1-a4["mean_ms"]/by["static",2]["mean_ms"])*100:.2f}%'+'，锁和原子任务池分别减少'+f'{(1-l4["mean_ms"]/by["lock",2]["mean_ms"])*100:.2f}%和{(1-o4["mean_ms"]/by["atomic",2]["mean_ms"])*100:.2f}%'+'。因此，仅给出一条“线程越多越快”的曲线会掩盖任务分配方式对扩展性的影响。')
para(f'4线程静态效率为{a4["efficiency"]*100:.2f}%，锁任务池为{l4["efficiency"]*100:.2f}%，原子任务池为{o4["efficiency"]*100:.2f}%，均未达到理想线性加速。线程管理、尾部等待、领取成本和共享资源会消耗收益；当前实验不具备把这些因素逐项分离的证据。')
para('部分2线程任务池配置的效率略高于100%，1线程锁版本也略快于串行。它们来自分别测得的均值，小幅超线性既可能涉及缓存状态，也可能受运行波动影响；不能将其解释为突破计算工作量约束。没有测量更多线程，不能声称4线程以上必然变慢。')
section(1)
heading('5.3 诊断证据：均分工作为何不等于同时完成',2)
import csv
rows=list(csv.DictReader((out/'线程诊断数据.csv').open()))
selected=sorted([r for r in rows if r['run']=='1' and r['threads']=='4'],key=lambda r:(['static','lock','atomic'].index(r['mode']),int(r['thread_id'])))
caption('表5 第1次4线程诊断（不是正式计时均值）','Table 5 First four-thread diagnostic run, separate from formal timing')
table(['方案','线程','任务块','内部行','领取次数','线程工作/ms'],[
[labels[r['mode']],r['thread_id'],r['tasks'],r['rows'],r['claims'],f'{float(r["thread_elapsed_ms"]):.3f}'] for r in selected])
pic('诊断_任务分配.png',15)
caption('图5 第1次4线程诊断的任务分配','Fig. 5 Task assignment in the first four-thread diagnostic run')
para('静态版本中，4个线程都完成16块，内部行数为1024、1024、1024、1022，但工作阶段耗时约为18.630、29.902、30.015、13.791 ms。行数仅差2行，无法从任务数量上解释如此大的时间差。并行区域仍需等待较慢线程结束，因此均分工作量没有消除尾部等待。')
para('锁任务池在同一次诊断中的任务数为21、11、11、21，线程工作时间约为19.654～19.902 ms；原子任务池也出现按需调整的分配。较早完成的线程领取更多块，使完成时刻更接近。这与正式测量中任务池优于静态的现象相符，但不是对所有硬件环境的因果证明。')
para('三次诊断的全部记录均已保留。第2次锁任务池诊断中仍存在明显时间差，说明动态分配不会消除所有干扰，也说明不能只选一幅诊断图概括稳定性能。诊断包含额外计数操作，其时间不可与表3、表4的正式均值直接拼接。')
section(2)
heading('5.4 环境约束与不能下的结论',2)
quota=info['cpu_stat'].split('---')
def counters(text):
 return {line.split()[0]:int(line.split()[1]) for line in text.splitlines() if len(line.split())==2}
before,after=map(counters,quota[:2])
para(f'正式计时窗口前后，cgroup的nr_throttled均为{before["nr_throttled"]}，throttled_usec均为{before["throttled_usec"]}，没有新增配额节流。配额用于确定最多4线程的主实验配置，但本次静态方案的性能差异不能直接归因于触发了容器限流。')
para('诊断能够观察到线程有效完成速率的差异，却未测量CPU频率、宿主机负载、硬件缓存事件或线程被暂停的具体原因。因此，不能确认“某个CPU较慢”或“内存带宽已饱和”。close绑定与cores位置保持固定，减少了一个配置变量，但不保证各工作线程拥有相同的独占物理资源。')
para('同样不能声称原子任务池在所有环境下逊于锁任务池。若要研究同步原语的纯成本，需要独立设计微基准；若要证明64行最优，需要改变任务粒度。两者均超出本次两项实验，不能用未经测量的推论补成实验结论。')
heading('5.5 图像结果与验收对应',2)
para('图6展示固定输入与三种4线程结果。矩形、圆和三角形轮廓均被提取；背景缓慢变化产生较弱响应。串行及9种并行输出的像素完全一致，分块接缝未引入额外边缘。该结论来自数值比较，缩略图仅便于阅读。')
heading('6 总结')
para('本文以数据依赖为起点，把Sobel内部行组织为固定任务块，并设计静态映射、锁任务池和原子任务池。锁保护真实共享的任务编号，计算过程依靠输出所有权避免互斥，并在阶段结束时统一同步。设计过程对应任务分解、组合、映射及共享内存同步，而不是单纯设置线程数量。')
para(f'本环境下，4线程锁任务池取得{l4["speedup"]:.3f}倍平均加速，静态分块取得{a4["speedup"]:.3f}倍，原子任务池取得{o4["speedup"]:.3f}倍。任务分配诊断表明，工作量均分不必然带来同时完成，动态领取可适应实际执行速率差异。结果同时保留了波动与诊断反例，没有把本次排序当作通用最优选择。')
para('局限包括单一合成图、单一虚拟环境、固定任务粒度，以及缺少硬件计数器归因。所有结论限定于预热后的核心计算，不扩展到自然图像数据集、端到端I/O、更多线程或其他算子。')
heading('7 参考文献')
refs=[
'[1] 袁平鹏. 并行程序设计：课程介绍[Z]. 华中科技大学，2026，课件页内编号58（考核要求）.',
'[2] 袁平鹏. 并行算法概述[Z]. 华中科技大学，2026，课件页内编号48—49、74—81（PCAM、粒度与映射）.',
'[3] 袁平鹏. 共享内存程序设计[Z]. 华中科技大学，2026，课件页内编号142、154—159、163（任务划分、同步与开发步骤）.',
'[4] 袁平鹏. 并行计算系统体系结构概述[Z]. 华中科技大学，2026，课件页内编号35—38（加速比、效率与并行开销）.',
'[5] Huamán A. Sobel derivatives[EB/OL]. OpenCV 4.10.0, 2024[2026-10-25]. https://github.com/opencv/opencv/blob/4.10.0/doc/tutorials/imgproc/imgtrans/sobel_derivatives/sobel_derivatives.markdown.',
'[6] Free Software Foundation. GNU Offloading and Multi Processing Runtime Library[EB/OL]. GCC 14.2.0, 2024[2026-10-25]. https://github.com/gcc-mirror/gcc/blob/releases/gcc-14.2.0/libgomp/libgomp.texi.'
]
for ref in refs:para(ref,False,size=9)
section(1)
pic('输入与输出对比.png',17)
caption('图6 固定输入及三种4线程输出（原始全尺寸图随附）','Fig. 6 Fixed input and four-thread outputs; full-size images are included')
heading('附录：复现与提交文件')
para('代码目录的README说明构建、单独运行、完整复现和报告生成方法。run_experiment.py生成固定输入、执行测试、保存原始300条计时与63条线程诊断记录，再绘制两项性能图表。重新测量会覆盖结果，因此必须同步更新报告。build_report.py直接读取本次汇总数据生成DOCX，不手工录入性能数值。')
para('实验结果目录保留输入PGM/PNG、串行与全部9种并行配置的输出PNG、原始及汇总CSV、环境元数据、源文件校验值、配额计数、验证日志和图表。二进制、中间PGM、旧版实验备份与原始模板放在仓库根目录temp中，不进入提交压缩包。')
para('学号、姓名及学院班级留空，待提交者填写。代码与报告分目录保存，提交包按“学号-姓名”命名。')
for p in d.paragraphs:
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for p in c.paragraphs:
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
props=d.core_properties
props.title='基于OpenMP的Sobel图像边缘检测：任务划分、同步机制与并行性能分析'
props.subject='并行程序设计课程作业'
props.author='';props.last_modified_by='';props.comments='';props.keywords='';props.category=''
props.identifier='';props.content_status='';props.revision=1
from docx.opc.constants import RELATIONSHIP_TYPE as RT
used={v for el in d._element.iter() for k,v in el.attrib.items() if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
for rid,rel in list(d.part.rels.items()):
 if rid not in used and (rel.reltype in [RT.IMAGE,RT.OLE_OBJECT] or any(x in rel.reltype.lower() for x in ['footnotes','endnotes','comments'])):
  d.part.drop_rel(rid)
report.mkdir(exist_ok=True)
from datetime import datetime, timezone
props.created=props.modified=datetime.now(timezone.utc)
for child in list(props._element):
 if child.tag.endswith('lastPrinted'):props._element.remove(child)
d.save(report/'程序设计报告.docx')
import zipfile
from lxml import etree
path=report/'程序设计报告.docx'
with zipfile.ZipFile(path) as z:
 items={name:z.read(name) for name in z.namelist()}
for name in ['docProps/custom.xml']:
 items.pop(name,None)
if 'docProps/app.xml' in items:
 app=etree.fromstring(items['docProps/app.xml'])
 for el in list(app):
  if etree.QName(el).localname in ['Company','Manager','Template','TotalTime','Pages','Words','Characters','Lines','Paragraphs','CharactersWithSpaces']:
   app.remove(el)
 items['docProps/app.xml']=etree.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
for name in ['_rels/.rels','[Content_Types].xml']:
 tree=etree.fromstring(items[name])
 for el in list(tree):
  if 'custom' in (el.get('Type','')+el.get('PartName','')).lower():tree.remove(el)
 items[name]=etree.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in items.items():z.writestr(name,data)

print(report/'程序设计报告.docx')
